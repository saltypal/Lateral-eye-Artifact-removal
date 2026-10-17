from pathlib import Path
import re
import build_full_report as base
from docx import Document
root=base.ROOT
text=(root/'docs/mentor/Full Research Report.md').read_text(encoding='utf-8')
parts=re.split(r'(?=^## \d+ )',text,flags=re.M)
selected={int(re.match(r'## (\d+)',p).group(1)):p for p in parts if re.match(r'## (\d+)',p)}
intro='''# Regional EOG Artifact Removal Experimental Results

Research results through 16 October 2026

## 1 Summary of measured outcomes

The latest controlled development comparison achieved 17.804727 dB whole-head channel-first mean output SNR and 18.261209 dB frontal mean output SNR using an anchored frontal VMD correction with a solver-failure raw fallback and matched nonfrontal regression. The direct-regression control achieved 17.785746 dB and 18.182199 dB, respectively. The incremental whole-head difference was +0.018981 dB with a supplemental 95% interval spanning zero. VMD superiority is therefore not established.

Native OSF evaluation showed reduced ocular association but excessive resting-signal modification. The classical teacher and EEG-only deployment model remain unqualified. This document presents measured results and their evaluation scope; historical branches use different corpora and must not be treated as one comparable leaderboard.

| Evidence set | Scope | Interpretation |
|---|---|---|
| Historical Klados replay | 24 windows from eight reused records; electrode identities unknown | Development-exposed EOG-assisted hybrid; not VMD alone |
| Latest controlled LEMON OSF corpus | 320 examples; 32 recipient windows; five ocular donors | Constructed contamination added to low-ocular reference EEG; not independent native ground truth |
| Latest native OSF pilot | Two development recordings; 51 scoring trials | Ocular suppression and resting preservation; no reconstruction SNR |
| Neural pilots | Five epochs on bounded development data | Preliminary learning results; no final model qualification |

Output SNR compares retained reference energy with reconstruction-error energy. Channel-first mean SNR and mean pooled-window SNR use different aggregation and are reported separately. Native EEG has no paired clean target and receives correlation and preservation measures. All numerical results were produced on Kaggle; this document assembles saved evidence.

<!-- PAGE -->

'''
# Keep measured experiment sections; remove presentation cues and theoretical subsections.
output=[intro]
for number in range(7,18):
    block=selected[number]
    block=re.sub(r'^## \d+ ',f'## {number-5} ',block,count=1)
    lines=block.splitlines()
    kept=[]
    skip=False
    for line in lines:
        if line.startswith('### Why TCN'):
            skip=True
        if skip:
            continue
        if line.startswith('**Say:**') or line.startswith('**Opening statement:**'):
            # Preserve evidence identifiers if they shared the speaking line.
            match=re.search(r'Evidence: (\[E\d+\].*)',line)
            if match: kept.append('Evidence: '+match.group(1))
            continue
        kept.append(line)
    output.append('\n'.join(kept))
output.append('''\n## 13 Evaluation interpretation and evidence access

Primary paired outcomes are MSE, temporal RRMSE, ordinary Pearson correlation, explicitly defined output SNR and SNR improvement. Native outcomes use ordinary absolute EEG HEOG correlation for lateral movement, EEG VEOG correlation for blink and vertical movement, resting RMSE and spectral preservation. EOG references remain unchanged between before and after calculations.

Project engineering limits are separate from published metric definitions: worst-source mean relative clean or resting modification at most 1%, alpha and beta power error at most 0.5 dB, relative covariance error at most 0.02, and paired correlation deterioration at most 0.005 against the declared comparator. The latest candidate fails preservation and is not production qualified.

The latest native run passed 341 numerical and source-contract tests and all 744 output payloads were recovered and hash verified. Test passage verifies implemented contracts, not scientific efficacy. No completed reserved confirmation, independent waveform benchmark or final EEG-only deployment qualification is claimed.

Evidence identifiers follow the companion Full Research Report and its evidence register. Source reports are retained in docs/FRONTAL_RESULTS.md, docs/GUARDED_REPLAY_RESULTS_20261011.md, docs/NATIVE_SEPARATE_AXIS_RESULTS.md, docs/FRONTAL_BOUNDED_FEASIBILITY_RESULTS.md, docs/FRONTAL_ARTIFACT_SHAPE_RESULTS.md, posterior results.md, docs/REGIONAL_RESULTS.md, docs/REGIONAL_DIAGNOSTIC_RESULTS.md, docs/NONFRONTAL_CHALLENGERS_RESULTS.md, docs/LATERAL_GAIN_RESULTS_20261015.md, docs/BROADER_FRONTAL_RESULTS_20261015.md and docs/NATIVE_FROZEN_FRONTAL_RESULTS_20261015.md. Evaluation equations and paper mappings are retained in docs/EVALUATION_PROTOCOL.md. Full run histories and detailed tables remain in Work So Far.md.

The latest native source commit is c48af6a0cda4d6bd905856613a3cf0c01f5bdcd2. Its artifact manifest SHA256 is ccc9253ee7b71f676bfacf805c5d7aabf3b4fed3cae147abad707c04ff071111.
''')
source=root/'docs/mentor/Experimental Results.md'
source.write_text('\n'.join(output),encoding='utf-8')
base.SOURCE=source
base.OUTPUT=root/'docs/mentor/Experimental Results.docx'
base.build()
doc=Document(base.OUTPUT)
for p in doc.paragraphs:
    if p.style.name=='Title': p.text='Regional EOG Artifact Removal'
    elif p.style.name=='Subtitle': p.text='Experimental results'
section=doc.sections[0]
section.header.paragraphs[0].text='REGIONAL EOG REMOVAL     EXPERIMENTAL RESULTS'
section.header.paragraphs[0].runs[0].font.size=base.Pt(8)
doc.core_properties.title='Regional EOG Artifact Removal Experimental Results'
doc.core_properties.subject='Measured research results and evaluation limitations'
doc.save(base.OUTPUT)
print('RESULTS_ONLY_READY',base.OUTPUT)
