"""Assemble a rehearsal report from saved evidence, without EEG computation."""
import hashlib
import json
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/mentor/Full Research Report.md"
OUTPUT = ROOT / "docs/mentor/Full Research Report.docx"
ASSETS = ROOT / "docs/mentor/report_assets"

FIGURES = {
    "modes": "vmd-broader-review-20261015/broader-frontal-review-001/results/vmd-broader-review-20261015/broader-frontal-review-001/predeclared_lateral_modes.png",
    "centers": "vmd-eog-20261010/vmd-001/results/vmd-eog-20261010/vmd-001/K_sweep_centers.png",
    "bounded": "vmd-frontal-replication-20261011/native-bounded-feasibility-001/results/vmd-frontal-replication-20261011/native-bounded-feasibility-001/bounded_native_preservation.png",
    "errors": "vmd-regional-challengers-20261013/regional-diagnostic-full-001/results/vmd-regional-challengers-20261013/regional-diagnostic-full-001/diagnostic_error_shares.png",
    "conditions": "vmd-broader-cached-20261015/broader-cached-review-001/results/vmd-broader-cached-20261015/broader-cached-review-001/broader_condition_snr.png",
}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rich_text(paragraph, text):
    """Keep URLs clickable and speaking cues bold, using native Word runs."""
    for token in re.split(r"(\*\*.*?\*\*|https?://[^\s]+)", text):
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith(("https://", "http://")):
            address = token.rstrip(".,;")
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(address,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
            run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "16528C")
            props.append(color)
            run.append(props)
            value = OxmlElement("w:t")
            labels = {"http://fcon_1000.projects.nitrc.org/indi/retro/MPI_LEMON/downloads/download_EEG.html": "LEMON EEG download page"}
            value.text = labels.get(address, address)
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
            paragraph.add_run(token[len(address):])
        else:
            paragraph.add_run(token)


def format_table(document, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    columns = len(rows[0])
    if any(len(row) != columns for row in rows):
        raise ValueError("Table column mismatch")
    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {2: (1.65, 5.05), 3: (2.15, 1.85, 2.70), 4: (2.65, 1.35, 1.35, 1.35),
              5: (2.50, 1.05, 1.05, 1.05, 1.05)}.get(columns)
    if columns == 2 and rows[0][0] == "ID":
        widths = (.50, 6.20)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    table._tbl.tblPr.append(borders)
    for row_number, values in enumerate(rows):
        row = table.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if row_number == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for column_number, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[column_number])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E9EEF3" if row_number == 0 else ("F7F8FA" if row_number % 2 == 0 else "FFFFFF"))
            properties.append(shade)
            margins = OxmlElement("w:tcMar")
            for edge, amount in (("top", 65), ("bottom", 65), ("left", 90), ("right", 90)):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), str(amount))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.02
            rich_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if row_number == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.70)
    section.left_margin = section.right_margin = Inches(.85)
    section.header_distance = section.footer_distance = Inches(.28)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.06
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 17), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    document.styles["Title"].font.size = Pt(24)
    document.styles["Title"].paragraph_format.space_after = Pt(5)
    document.styles["Subtitle"].font.size = Pt(10)
    document.styles["Subtitle"].font.italic = False
    for name in ("Title", "Subtitle"):
        rpr = document.styles[name]._element.get_or_add_rPr()
        for node in list(rpr.findall(qn("w:spacing"))):
            rpr.remove(node)
    section.header.paragraphs[0].add_run("REGIONAL EOG REMOVAL     RESEARCH AND REHEARSAL").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("16 October 2026     |     Page ").font.size = Pt(8)
    for instruction in ("PAGE", "NUMPAGES"):
        if instruction == "NUMPAGES":
            footer.add_run(" of ").font.size = Pt(8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        footer._p.append(field)
    ASSETS.mkdir(parents=True, exist_ok=True)
    figure_register = {}
    for key, relative in FIGURES.items():
        source = ROOT / "results" / relative
        manifest = json.loads((source.parent / "artifact_manifest.json").read_text(encoding="utf-8"))
        digest = sha256(source)
        if manifest.get(source.name) != digest:
            raise ValueError("Figure differs from original run manifest: " + key)
        destination = ASSETS / (key + ".png")
        shutil.copyfile(source, destination)
        figure_register[key] = {"source": str(source), "sha256": digest,
                                "signal_processing": False, "new_numerical_plot": False}
    pending_break = False
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- PAGE -->":
            pending_break = True
            index += 1
            continue
        if line.startswith("# "):
            document.add_paragraph("Regional EOG Artifact Removal", style="Title")
            document.add_paragraph("Research and rehearsal report", style="Subtitle")
        elif line.startswith("## "):
            heading = document.add_paragraph(line[3:], style="Heading 1")
            heading.paragraph_format.page_break_before = pending_break
            pending_break = False
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith(":::figure "):
            _, key, width, height = line.split()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(ASSETS / (key + ".png")),
                                            width=Inches(float(width)), height=Inches(float(height)))
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            format_table(document, table_lines)
            continue
        else:
            paragraph = document.add_paragraph()
            rich_text(paragraph, line)
            if line.startswith("**Figure"):
                paragraph.paragraph_format.space_after = Pt(7)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if line.startswith("Research record"):
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        index += 1
    document.core_properties.title = "Regional EOG Artifact Removal Research and Rehearsal Report"
    document.core_properties.subject = "Methods evidence failed hypotheses and mentor presentation preparation"
    document.core_properties.comments = "Development research; no production qualification or native reconstruction SNR claim."
    document.save(OUTPUT)
    (ASSETS / "figure_provenance.json").write_text(json.dumps(figure_register, indent=2), encoding="utf-8")
    print(OUTPUT)
    print("EXISTING_FIGURES_VERIFIED", len(figure_register))


if __name__ == "__main__":
    build()
