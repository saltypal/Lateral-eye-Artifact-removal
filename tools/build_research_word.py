"""Convert the verified report to editable Word with embedded original figures."""
from pathlib import Path
import base64
import io
import re

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1] / 'docs' / 'research'


def inline(paragraph, node, bold=False, italic=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold, run.italic = bold, italic
        return
    if node.name == 'a':
        link = OxmlElement('w:hyperlink')
        relation = paragraph.part.relate_to(node.get('href', ''),
            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        link.set(qn('r:id'), relation)
        run = OxmlElement('w:r')
        properties = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '145A86')
        properties.append(color); run.append(properties)
        text = OxmlElement('w:t'); text.text = node.get_text(); run.append(text)
        link.append(run); paragraph._p.append(link)
        return
    for child in node.children:
        inline(paragraph, child, bold or node.name in ('strong', 'b'), italic or node.name in ('em', 'i'))


def build():
    soup = BeautifulSoup((ROOT / 'Methods and Results.html').read_text(encoding='utf-8'), 'html.parser')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.7)
    section.top_margin = section.bottom_margin = Inches(.65)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    images = 0
    for node in soup.body.children:
        if not getattr(node, 'name', None):
            continue
        if 'FIGURE_ATLAS' in node.get_text():
            continue
        if node.name in ('h1', 'h2', 'h3'):
            label = re.sub(r'[^\w\s]', ' ', node.get_text()).strip()
            doc.add_paragraph(label, 'Title' if node.name == 'h1' else 'Heading ' + str(int(node.name[1]) - 1))
        elif node.name == 'table':
            rows = node.find_all('tr')
            count = len(rows[0].find_all(['th', 'td']))
            table = doc.add_table(rows=0, cols=count)
            table.style = 'Table Grid'
            widths = [2.3] + [(7.1 - 2.3) / (count - 1)] * (count - 1)
            if count == 3:
                widths = [1.65, 2.55, 2.9]
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for column, content in enumerate(row.find_all(['th', 'td'])):
                    cells[column].width = Inches(widths[column])
                    p = cells[column].paragraphs[0]
                    p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.space_before = Pt(5)
                    for child in content.children:
                        inline(p, child, bold=index == 0)
                    for run in p.runs:
                        run.font.size = Pt(9)
                    if index == 0:
                        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E5EDF2')
                        cells[column]._tc.get_or_add_tcPr().append(shading)
                if index == 0:
                    header = OxmlElement('w:tblHeader'); table.rows[-1]._tr.get_or_add_trPr().append(header)
            doc.add_paragraph()
        elif node.name == 'p' and node.find('img'):
            data = base64.b64decode(node.img['src'].split(',', 1)[1])
            with Image.open(io.BytesIO(data)) as image:
                ratio = image.height / image.width
            width = min(7.1, 7.8 / ratio)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(io.BytesIO(data), width=Inches(width))
            images += 1
        elif node.name in ('p', 'pre', 'ul', 'ol'):
            p = doc.add_paragraph()
            for child in node.children:
                inline(p, child)
            if node.name == 'pre':
                for run in p.runs:
                    run.font.name, run.font.size = 'Consolas', Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    output = ROOT / 'Methods and Results.docx'
    doc.save(output)
    assert images == 54 and len(doc.inline_shapes) == 54
    print(f'Saved {output}; {images} embedded figures; {len(doc.tables)} editable tables')


if __name__ == '__main__':
    build()
