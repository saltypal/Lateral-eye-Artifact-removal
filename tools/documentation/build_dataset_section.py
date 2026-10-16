"""Build the first mentor-review section without processing research signals."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/mentor/sections/01_datasets.md"
OUTPUT = ROOT / "docs/mentor/Research Documentation.docx"


def linked_text(paragraph, text):
    position = 0
    for match in re.finditer(r"https?://[^\s]+", text):
        paragraph.add_run(text[position:match.start()])
        address = match.group().rstrip(".")
        link = OxmlElement("w:hyperlink")
        relationship = paragraph.part.relate_to(
            address, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link.set(qn("r:id"), relationship)
        run, properties = OxmlElement("w:r"), OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "16528C")
        properties.append(color)
        run.append(properties)
        value = OxmlElement("w:t")
        labels = {
            "http://fcon_1000.projects.nitrc.org/indi/retro/MPI_LEMON/downloads/download_EEG.html": "LEMON EEG downloads",
            "https://ftp.gwdg.de/pub/misc/MPI-Leipzig_Mind-Brain-Body-LEMON/": "MPI Leipzig Mind Brain Body archive",
        }
        value.text = labels.get(address, address)
        run.append(value)
        link.append(run)
        paragraph._p.append(link)
        if match.group().endswith("."):
            paragraph.add_run(".")
        position = match.end()
    paragraph.add_run(text[position:])


def build():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    section.header_distance = section.footer_distance = Inches(.3)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    document.styles["Subtitle"].font.italic = False
    for name in ("Title", "Subtitle"):
        properties = document.styles[name]._element.get_or_add_rPr()
        spacing = properties.find(qn("w:spacing"))
        if spacing is not None:
            properties.remove(spacing)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    document.styles["Title"].font.size = Pt(23)
    document.styles["Title"].paragraph_format.space_after = Pt(4)
    document.styles["Subtitle"].font.size = Pt(11)
    document.styles["Subtitle"].paragraph_format.space_after = Pt(14)
    document.add_paragraph("EEG Ocular Artifact Removal", style="Title")
    document.add_paragraph("Research documentation  |  Dataset section  |  16 October 2026", style="Subtitle")
    header = section.header.paragraphs[0]
    header.add_run("MENTOR REVIEW     DATASETS AND SOURCES").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for block in SOURCE.read_text(encoding="utf-8").strip().split("\n\n"):
        if block.startswith("# "):
            document.add_paragraph("1 " + block[2:], style="Heading 1")
        elif block.startswith("## "):
            heading = document.add_paragraph(block[3:], style="Heading 2")
            if block[3:] == "References":
                heading.paragraph_format.page_break_before = True
        else:
            paragraph = document.add_paragraph()
            linked_text(paragraph, block)
            if block.startswith(tuple(f"[{i}]" for i in range(1, 7))):
                paragraph.paragraph_format.space_after = Pt(5)
    document.core_properties.title = "EEG Ocular Artifact Removal Research Documentation"
    document.core_properties.subject = "Datasets used and source references"
    document.core_properties.comments = "Dataset section only; subsequent sections await user review."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
